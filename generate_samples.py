import openpyxl
from openpyxl.styles import Font, PatternFill
import os

os.makedirs("samples", exist_ok=True)

# Candidates sample
wb_c = openpyxl.Workbook()
ws_c = wb_c.active
ws_c.title = "Candidates"
ws_c.append(["Name", "Mobile", "Registration ID"])
for col in range(1, 4):
    cell = ws_c.cell(row=1, column=col)
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

data_c = [
    ["Rahul Sharma", "9876543210", "TH-2026-001"],
    ["Priya Patel", "8765432109", "TH-2026-002"],
    ["Arjun Kumar", "7654321098", "TH-2026-003"],
    ["Sneha Reddy", "9988776655", "TH-2026-004"],
    ["Vikram Singh", "9123456789", "TH-2026-005"],
    ["Anjali Mehta", "9871234567", "TH-2026-006"],
]
for row in data_c:
    ws_c.append(row)

wb_c.save("samples/candidates_sample.xlsx")
print("Created samples/candidates_sample.xlsx with 6 dummy candidates")

# Questions sample
wb_q = openpyxl.Workbook()
ws_q = wb_q.active
ws_q.title = "Questions"
ws_q.append(["Question", "OptionA", "OptionB", "OptionC", "OptionD", "Answer"])
for col in range(1, 7):
    cell = ws_q.cell(row=1, column=col)
    cell.font = Font(bold=True)
    cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

data_q = [
    ["What is the capital of France?", "London", "Berlin", "Paris", "Madrid", "C"],
    ["Which keyword is used to define a function in Python?", "func", "def", "function", "define", "B"],
    ["What does HTML stand for?", "Hyper Text Markup Language", "High Transfer Machine Language", "Hyperlink Text Management", "None of these", "A"],
    ["What is 2 + 2 * 3?", "12", "8", "10", "18", "B"],
    ["Which of these is mutable in Python?", "Tuple", "String", "List", "Integer", "C"],
    ["What is the output of print(2**3)?", "6", "8", "9", "5", "B"],
    ["Django is a Python web framework: True or False?", "False", "True", "Maybe", "Error", "B"],
    ["What does SQL stand for?", "Structured Query Language", "Simple Query List", "System Query Language", "None", "A"],
]
for row in data_q:
    ws_q.append(row)

wb_q.save("samples/questions_sample.xlsx")
print("Created samples/questions_sample.xlsx with 8 dummy questions")
print("\n=== Dummy files created successfully in samples/ folder ===")
print("Use candidates_sample.xlsx for Bulk Import Candidates")
print("Use questions_sample.xlsx for Bulk Import Questions (can be scoped to an exam)")
print("\nFormat notes:")
print("- Candidates: Name, Mobile, Registration ID (password auto-generated)")
print("- Questions: Question, OptionA, OptionB, OptionC, OptionD, Answer (use A/B/C/D)")
print("\nThe bulk import functions in views.py are already set up and ready.")
